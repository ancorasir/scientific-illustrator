"""Synthetic OOXML round-trip tests; no grant documents or native UI claims."""
import importlib.util
import json
from pathlib import Path
import tempfile
from zipfile import ZipFile
from xml.etree import ElementTree as E
from pptx import Presentation
from pptx.util import Pt

root = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("bridge", root / "plugins/scientific-illustrator/scripts/powerpoint-mac-bridge.py")
bridge = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bridge)
deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[6])
shape = slide.shapes.add_textbox(Pt(20),Pt(20),Pt(150),Pt(60))
shape.name = "synthetic-method"
shape.text = "示例方法\nGeneric method"
untouched = slide.shapes.add_textbox(Pt(20),Pt(100),Pt(150),Pt(30))
untouched.text = "Preserve this object"
before = untouched._element.xml
args = {"font_name":"MiSans","font_size":10,"line_spacing_pt":12.5,
        "paragraph_space_before_pt":2,"paragraph_space_after_pt":3,"text_autofit":"none"}
bridge._apply_text_frame(shape.text_frame,args)
assert untouched._element.xml == before
n = {"a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
with tempfile.TemporaryDirectory(prefix="figure-typography-test-") as d:
    p = Path(d)/"synthetic.pptx"
    deck.save(p)
    reopened = Presentation(p)
    assert reopened.slides[0].shapes[0].text == shape.text
    with ZipFile(p) as z:
        q = E.fromstring(z.read("ppt/slides/slide1.xml"))
    for ppr in q.findall(".//a:pPr",n):
        # Only the explicitly styled shape has paragraph settings.
        if ppr.find("a:lnSpc",n) is None: continue
        assert ppr.find("a:lnSpc/a:spcPts",n).get("val") == "1250"
        assert ppr.find("a:spcBef/a:spcPts",n).get("val") == "200"
        assert ppr.find("a:spcAft/a:spcPts",n).get("val") == "300"
    assert len(q.findall(".//a:lnSpc/a:spcPts",n)) == 2
    for tag in ["latin","ea","cs"]:
        assert {v.get("typeface") for v in q.findall(f".//a:rPr/a:{tag}",n)} == {"MiSans"}
    assert not q.findall(".//a:normAutofit",n)
for field,value in [("line_spacing_pt",0),("line_spacing_pt",float("nan")),("paragraph_space_after_pt",-1)]:
    try: bridge._apply_text_frame(shape.text_frame,{field:value})
    except ValueError: pass
    else: raise AssertionError(f"Invalid {field} accepted")
# Omitting paragraph settings must retain explicit values on a later font edit.
bridge._apply_text_frame(shape.text_frame,{"font_size":11})
assert shape.text_frame.paragraphs[0].line_spacing.pt == 12.5
print("OOXML explicit line/paragraph spacing, CJK font fields, untouched-object preservation and save/reopen passed (no native UI certification).")
