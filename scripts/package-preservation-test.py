"""Synthetic-only regression fixtures. No real documents or network dependencies."""
import copy
import importlib.util
import io
import json
from pathlib import Path
import random
import tempfile
import unittest
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("package_tools", ROOT/"plugins/scientific-illustrator/scripts/pptx-package.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)


def fixture(path):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
    a = slide.shapes.add_group_shape(); a.name = "Group A"
    s = a.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(20), Pt(30), Pt(70), Pt(40)); s.name = "Input"
    p = s.text_frame.paragraphs[0]
    for value, bold in [("Alpha ",False),("Beta",True)]:
        run=p.add_run();run.text=value;run.font.name="MiSans";run.font.size=Pt(12);run.font.bold=bold
    peer=a.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(90),Pt(100),Pt(60),Pt(40));peer.name="Peer";peer.text="Peer node"
    nested=a.shapes.add_group_shape();nested.name="Nested"
    n=nested.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(130),Pt(40),Pt(50),Pt(35));n.name="Nested target";n.text="Nested node"
    b=slide.shapes.add_group_shape();b.name="Group B"
    t=b.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(270),Pt(100),Pt(65),Pt(40));t.name="Output";t.text="Result"
    for parent, source, target, name in [(a.shapes,s,peer,"Internal"),(slide.shapes,n,t,"Cross"),(slide.shapes,s,a,"To group")]:
        c=parent.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(1),Pt(1),Pt(20),Pt(20));c.name=name
        c.begin_connect(source,3);c.end_connect(target,1)
    slide.notes_slide.notes_text_frame.text="Synthetic note only"
    second=prs.slides.add_slide(prs.slide_layouts[6]);second.shapes.add_textbox(Pt(20),Pt(20),Pt(100),Pt(50)).text="Unselected slide"
    prs.save(path)
    # Unknown members must survive too; the patcher must not discard them.
    with ZipFile(path,"a") as z:z.writestr("custom/synthetic.dat",b"opaque synthetic bytes")
    return path


def rewrite(path, fn):
    with ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    root=m.xml(parts["ppt/slides/slide1.xml"]);fn(root)
    parts["ppt/slides/slide1.xml"]=E.tostring(root)
    with ZipFile(path,"w") as z:
        for n,data in parts.items():z.writestr(n,data)


class PackageTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.path=fixture(Path(self.tmp.name)/"fixture.pptx");self.out=Path(self.tmp.name)/"candidate.pptx"
    def plan(self, **fields):
        return m.plan_local_patch(self.path,[{"slide_index":1,"shape_name":"Input",**fields}])
    def apply(self, plan, **overrides):
        args=dict(source_path=self.path,output_path=self.out,expected_source_sha256=plan["source_sha256"],edits=plan["edits"]);args.update(overrides)
        return m.apply_local_patch(**args)
    def test_recursive_group_and_cross_group_refs(self):
        report=m.audit_connections(self.path,include_objects=True);row=report["slides"][0]
        self.assertTrue(report["passed"]);self.assertEqual(row["group_count"],3);self.assertEqual(row["shape_count"],4)
        self.assertEqual(row["connector_count"],3);self.assertEqual(row["valid_connector_count"],3)
        self.assertEqual(row["cross_group_connector_count"],2);self.assertEqual(row["group_internal_connector_count"],1)
        self.assertTrue(all(o["world_bounds_pt"] is not None for o in row["objects"]))
    def test_missing_reference_and_site_failure(self):
        rewrite(self.path,lambda r:r.find('.//a:stCxn',m.NS).set('id','99999'))
        self.assertFalse(m.audit_connections(self.path)["passed"])
        self.path=fixture(self.path)
        rewrite(self.path,lambda r:r.find('.//a:stCxn',m.NS).set('idx','4'))
        self.assertIn('invalid_connection_site',json.dumps(m.audit_connections(self.path)))
    def test_unresolved_site_not_guessed(self):
        rewrite(self.path,lambda r:r.find('.//a:prstGeom',m.NS).set('prst','hexagon'))
        result=m.audit_connections(self.path)
        self.assertTrue(result['passed']);self.assertIn('connection_site_count_unresolved',json.dumps(result))
    def test_custom_geometry_site_count(self):
        def change(r):
            sp=r.find('.//p:spPr',m.NS);geom=sp.find('a:prstGeom',m.NS);sp.remove(geom)
            custom=E.SubElement(sp,'{'+m.NS['a']+'}custGeom');E.SubElement(custom,'{'+m.NS['a']+'}cxnLst')
        rewrite(self.path,change)
        self.assertFalse(m.audit_connections(self.path)['passed'])
    def test_duplicate_identity(self):
        def change(r):
            items=m.records(r);items[1]['element'].find('.//p:cNvPr',m.NS).set('id',items[0]['id'])
        rewrite(self.path,change)
        self.assertFalse(m.audit_connections(self.path)['passed'])
        with self.assertRaises(ValueError):self.plan(text='Changed')
    def test_run_style_and_opaque_parts_preserved(self):
        before=m.Package(self.path);plan=self.plan(text='Alpha better Beta');receipt=self.apply(plan)
        after=m.Package(self.out);self.assertTrue(receipt['preservation_passed'])
        self.assertEqual(receipt['changed_parts'],['ppt/slides/slide1.xml'])
        for n,data in before.parts.items():
            if n not in receipt['changed_parts']:self.assertEqual(data,after.parts[n],n)
        old={v['name']:v for v in m.records(before.slide(1)[1])};new={v['name']:v for v in m.records(after.slide(1)[1])}
        self.assertEqual(m.masked(old['Input']['element'],{'text':''}),m.masked(new['Input']['element'],{'text':''}))
        self.assertEqual(len(new['Input']['element'].findall('p:txBody/a:p/a:r',m.NS)),2)
        self.assertEqual(''.join(n.text or '' for n in new['Input']['element'].findall('p:txBody//a:t',m.NS)),'Alpha better Beta')
        for name in old:
            if name!='Input':self.assertEqual(old[name]['shallow_sha256'],new[name]['shallow_sha256'])
        Presentation(self.out)
    def test_random_run_edit_order(self):
        random.seed(31)
        for _ in range(120):
            old=''.join(random.choice('abc') for _ in range(random.randint(0,20)))
            text=''.join(random.choice('abcd') for _ in range(random.randint(0,25)))
            split=random.randint(0,len(old));nodes=[E.Element('t'),E.Element('t')];nodes[0].text=old[:split];nodes[1].text=old[split:]
            m.replace_segment(nodes,text);self.assertEqual(''.join(n.text or '' for n in nodes),text)
    def test_local_grouped_geometry_and_binding_preservation(self):
        with self.assertRaisesRegex(ValueError,'coordinate_space'):self.plan(left=25)
        plan=self.plan(left=25,width=74,coordinate_space='local');receipt=self.apply(plan)
        self.assertTrue(receipt['preservation_passed']);self.assertTrue(m.audit_connections(self.out)['passed'])
        self.assertEqual(m.audit_connections(self.path)['slides'][0]['connectors'],m.audit_connections(self.out)['slides'][0]['connectors'])
    def test_group_movement(self):
        plan=m.plan_local_patch(self.path,[{'slide_index':1,'shape_name':'Group A','left':50}]);self.apply(plan)
        before=m.audit_connections(self.path,include_objects=True)['slides'][0]['objects'];after=m.audit_connections(self.out,include_objects=True)['slides'][0]['objects']
        old={o['name']:o for o in before};new={o['name']:o for o in after}
        self.assertEqual(new['Input']['local_bounds_pt'],old['Input']['local_bounds_pt'])
        self.assertNotEqual(new['Input']['world_bounds_pt'],old['Input']['world_bounds_pt'])
    def test_rotation_flip_and_scale(self):
        def change(r):
            group=next(o for o in m.records(r) if o['name']=='Group A')['element']
            xf=group.find('p:grpSpPr/a:xfrm',m.NS);xf.set('rot','5400000');xf.set('flipH','1')
            ex=xf.find('a:ext',m.NS);ex.set('cx',str(2*int(ex.get('cx'))))
        rewrite(self.path,change)
        obj=next(o for o in m.audit_connections(self.path,include_objects=True)['slides'][0]['objects'] if o['name']=='Input')
        self.assertAlmostEqual(obj['world_bounds_pt'][2],40);self.assertAlmostEqual(obj['world_bounds_pt'][3],140)
    def test_bad_group_transform_warns(self):
        rewrite(self.path,lambda r:r.find('.//a:chExt',m.NS).set('cx','0'))
        result=m.audit_connections(self.path);self.assertTrue(result['passed']);self.assertIn('world_geometry_unresolved',json.dumps(result))
    def test_stale_source_rejected(self):
        plan=self.plan(text='Changed');rewrite(self.path,lambda r:r.find('.//a:t',m.NS).__setattr__('text','User update'))
        current=self.path.read_bytes()
        with self.assertRaisesRegex(ValueError,'Source conflict'):self.apply(plan)
        self.assertEqual(current,self.path.read_bytes());self.assertFalse(self.out.exists())
    def test_object_precondition_rejected(self):
        with self.assertRaisesRegex(ValueError,'Object conflict'):self.plan(text='Changed',expected_shape_sha256='0'*64)
    def test_no_overwrite(self):
        plan=self.plan(text='Changed')
        with self.assertRaises(ValueError):self.apply(plan,output_path=self.path)
        self.out.write_bytes(b'keep')
        with self.assertRaises(ValueError):self.apply(plan)
        self.assertEqual(self.out.read_bytes(),b'keep')
    def test_paragraph_and_dynamic_fields_rejected(self):
        with self.assertRaisesRegex(ValueError,'paragraph'):self.plan(text='First\nSecond')
        rewrite(self.path,lambda r:E.SubElement(r.find('.//p:txBody/a:p',m.NS),'{'+m.NS['a']+'}fld'))
        with self.assertRaisesRegex(ValueError,'dynamic'):self.plan(text='Changed')
    def test_ambiguous_name(self):
        rewrite(self.path,lambda r:next(o for o in m.records(r) if o['name']=='Output')['element'].find('.//p:cNvPr',m.NS).set('name','Input'))
        with self.assertRaisesRegex(ValueError,'ambiguous'):self.plan(text='Changed')
        report=m.audit_connections(self.path,include_objects=True);shape=next(o for o in report['slides'][0]['objects'] if o['name']=='Input')
        plan=m.plan_local_patch(self.path,[{'slide_index':1,'shape_id':int(shape['id']),'text':'Changed'}]);self.apply(plan)
    def test_invalid_edits_are_atomic(self):
        original=self.path.read_bytes()
        bad=[{'slide_index':1,'shape_name':'Input','text':'Changed'},{'slide_index':1,'shape_name':'Output','width':-1}]
        with self.assertRaises(ValueError):m.apply_local_patch(self.path,self.out,m.digest(original),bad)
        self.assertFalse(self.out.exists());self.assertEqual(original,self.path.read_bytes())
    def test_ancestor_descendant_overlap(self):
        for edits in [[{'slide_index':1,'shape_name':'Group A','left':20},{'slide_index':1,'shape_name':'Input','text':'Changed'}], [{'slide_index':1,'shape_name':'Input','text':'Changed'},{'slide_index':1,'shape_name':'Group A','left':20}]]:
            with self.assertRaisesRegex(ValueError,'Ancestor'):m.plan_local_patch(self.path,edits)
    def test_invalid_fields_and_values(self):
        for fields in [{'font_size':12},{'left':float('nan')},{'left':True},{'width':0},{'text':123}]:
            with self.assertRaises(ValueError):self.plan(**fields)
    def test_noop_plan_and_slide_bounds(self):
        self.assertEqual(self.plan(text='Alpha Beta')['changed_parts'],[])
        with self.assertRaises(ValueError):m.audit_connections(self.path,[])
        for index in [0,3,True]:
            with self.assertRaises(ValueError):m.audit_connections(self.path,[index])
    def test_signed_package_rejected(self):
        with ZipFile(self.path,'a') as z:z.writestr('_xmlsignatures/sig1.xml',b'synthetic')
        with self.assertRaisesRegex(ValueError,'Signed'):self.plan(text='Changed')
    def test_readonly_plan_and_late_conflict(self):
        original=self.path.read_bytes();self.plan(text='Changed');self.assertEqual(self.path.read_bytes(),original)
        from unittest.mock import patch
        prepare=m.prepare_patch
        def concurrent(package,edits):
            result=prepare(package,edits)
            rewrite(self.path,lambda r:r.find('.//a:t',m.NS).__setattr__('text','Concurrent user edit'))
            return result
        plan=self.plan(text='Changed')
        with patch.object(m,'prepare_patch',side_effect=concurrent):
            with self.assertRaisesRegex(ValueError,'during patch preparation'):self.apply(plan)
        self.assertFalse(self.out.exists())
    def test_existing_figure_audit_reads_imported_bindings(self):
        from unittest.mock import patch
        spec=importlib.util.spec_from_file_location('bridge',ROOT/'plugins/scientific-illustrator/scripts/powerpoint-mac-bridge.py')
        bridge=importlib.util.module_from_spec(spec);spec.loader.exec_module(bridge)
        with patch.object(bridge,'_load',return_value=({'metadata':{}},self.path,Presentation(self.path))):
            result=bridge.action_audit_figure({'slide_index':1})
        self.assertTrue(result['connections']['passed'])
        self.assertEqual(result['connections']['slides'][0]['connector_count'],3)
    def test_soft_break_structure_and_styles_preserved(self):
        def change(r):
            paragraph=r.find('.//p:txBody/a:p',m.NS)
            br=E.Element('{'+m.NS['a']+'}br');paragraph.insert(paragraph.index(paragraph.findall('a:r',m.NS)[1]),br)
        rewrite(self.path,change)
        plan=self.plan(text='Alpha revised\nBeta');self.apply(plan)
        before=next(o for o in m.records(m.Package(self.path).slide(1)[1]) if o['name']=='Input')
        after=next(o for o in m.records(m.Package(self.out).slide(1)[1]) if o['name']=='Input')
        self.assertEqual(m.masked(before['element'],{'text':''}),m.masked(after['element'],{'text':''}))
    def test_macro_rejected(self):
        with ZipFile(self.path,'a') as z:z.writestr('ppt/vbaProject.bin',b'synthetic')
        with self.assertRaisesRegex(ValueError,'macro'):self.plan(text='Changed')
    def test_presentation_order(self):
        with ZipFile(self.path) as z:parts={n:z.read(n) for n in z.namelist()}
        root=m.xml(parts['ppt/presentation.xml']);lst=root.find('p:sldIdLst',m.NS);lst[:]=list(reversed(lst))
        parts['ppt/presentation.xml']=E.tostring(root)
        with ZipFile(self.path,'w') as z:
            for n,data in parts.items():z.writestr(n,data)
        self.assertEqual(m.audit_connections(self.path)['slides'][0]['connector_count'],0)
        self.assertEqual(m.audit_connections(self.path)['slides'][1]['connector_count'],3)

if __name__=='__main__':
    import sys
    if len(sys.argv)>1 and sys.argv[1]=='--fixture':
        fixture(Path(sys.argv[2]))
    else:
        unittest.main(verbosity=2)
